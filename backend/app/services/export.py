"""Excel export of the full history (TZ §11 / §14 premium)."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Category, Transaction


async def export_xlsx(session: AsyncSession, user_id: int, base_currency: str, lang: str) -> bytes:
    rows = (await session.execute(
        select(Transaction, Category.name)
        .outerjoin(Category, Category.id == Transaction.category_id)
        .where(Transaction.user_id == user_id)
        .order_by(Transaction.occurred_at.desc(), Transaction.id.desc())
    )).all()

    ru = lang == "ru"
    headers = (["Дата", "Тип", "Категория", "Описание", f"Сумма ({base_currency})", "Исходно", "Валюта", "Источник"]
               if ru else
               ["Date", "Type", "Category", "Title", f"Amount ({base_currency})", "Original", "Currency", "Source"])

    wb = Workbook()
    ws = wb.active
    ws.title = "NEBANK" if not ru else "NEBANK"
    head_fill = PatternFill("solid", fgColor="F2F4F8")
    head_font = Font(bold=True, color="1A1A1A")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="left")

    for i, (tx, cat) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=tx.occurred_at.isoformat())
        ws.cell(row=i, column=2, value=("Доход" if ru else "Income") if tx.kind.value == "income"
                else ("Расход" if ru else "Expense"))
        ws.cell(row=i, column=3, value=cat or "")
        ws.cell(row=i, column=4, value=tx.title or "")
        ws.cell(row=i, column=5, value=float(tx.amount))
        ws.cell(row=i, column=6, value=float(tx.original_amount))
        ws.cell(row=i, column=7, value=tx.original_currency)
        ws.cell(row=i, column=8, value=tx.source.value)

    widths = [12, 10, 20, 28, 16, 14, 10, 10]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
